"""P0-6: evidence uploads validate real content type (magic bytes), not just extension."""
import io
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, SimpleTestCase, override_settings
from django.urls import reverse

from compliance.upload_validation import validate_evidence_file
from compliance.models import Evidence, EvidenceSubmission
from compliance.tests import _company_with_control, _journey_user, _company_with_submission

ALLOWED = ['pdf', 'png', 'jpg', 'jpeg', 'tiff', 'bmp', 'docx', 'xlsx', 'txt']
PNG = bytes.fromhex('89504e470d0a1a0a') + b'\x00' * 32
PDF = b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n'
EXE = b'MZ\x90\x00' + b'\x00' * 40


def _docx_bytes():
    # A REAL python-docx file (multi-entry OOXML) — its 'word/' marker sits well past a short
    # header, which is exactly what broke real .docx uploads (MAT3 DEF-02). A toy 2-entry zip
    # would fit in 262 bytes and give a false pass, so use the real writer.
    from docx import Document
    b = io.BytesIO()
    d = Document()
    d.add_paragraph('Approved cybersecurity policy. ' * 40)
    d.save(b)
    return b.getvalue()


def _xlsx_bytes():
    from openpyxl import Workbook
    b = io.BytesIO()
    Workbook().save(b)
    return b.getvalue()


def _f(name, data):
    return SimpleUploadedFile(name, data)


class ValidateEvidenceFileUnitTests(SimpleTestCase):
    def test_real_pdf_ok(self):
        ok, ext, _ = validate_evidence_file(_f('policy.pdf', PDF), ALLOWED)
        self.assertTrue(ok); self.assertEqual(ext, 'pdf')

    def test_exe_renamed_pdf_rejected(self):
        ok, _, err = validate_evidence_file(_f('malware.pdf', EXE), ALLOWED)
        self.assertFalse(ok); self.assertIn('does not match', err)

    def test_png_renamed_jpg_rejected(self):
        ok, _, _ = validate_evidence_file(_f('img.jpg', PNG), ALLOWED)
        self.assertFalse(ok)                      # content (png) != declared (jpg)

    def test_real_png_ok(self):
        ok, _, _ = validate_evidence_file(_f('img.png', PNG), ALLOWED)
        self.assertTrue(ok)

    def test_real_docx_ok(self):
        ok, ext, err = validate_evidence_file(_f('doc.docx', _docx_bytes()), ALLOWED)
        self.assertTrue(ok, err); self.assertEqual(ext, 'docx')   # MAT3 DEF-02 regression guard

    def test_real_xlsx_ok(self):
        ok, ext, err = validate_evidence_file(_f('book.xlsx', _xlsx_bytes()), ALLOWED)
        self.assertTrue(ok, err); self.assertEqual(ext, 'xlsx')

    def test_text_file_no_magic_ok(self):
        ok, _, _ = validate_evidence_file(_f('note.txt', b'just plain policy text'), ALLOWED)
        self.assertTrue(ok)                       # None sniff + txt in text allowlist

    def test_text_content_named_pdf_rejected(self):
        ok, _, _ = validate_evidence_file(_f('fake.pdf', b'not a real pdf, plain text'), ALLOWED)
        self.assertFalse(ok)                      # None sniff + pdf not a text type

    def test_disallowed_extension_rejected(self):
        ok, _, _ = validate_evidence_file(_f('x.exe', EXE), ALLOWED)
        self.assertFalse(ok)


class UploadViewRejectionTests(TestCase):
    def test_control_detail_upload_rejects_spoofed_file(self):
        c, control = _company_with_control()
        self.client.force_login(_journey_user(c, email='p6a@x.com'))
        resp = self.client.post(reverse('compliance:upload_evidence', args=[control.id]),
                                {'evidence_file': _f('malware.pdf', EXE)})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Evidence.objects.count(), 0)   # spoofed file never stored

    def test_checklist_v2_upload_rejects_spoofed_file(self):
        c, item, sub = _company_with_submission(fv_code='NCA-ECC-2-2024')
        before = EvidenceSubmission.objects.count()
        self.client.force_login(_journey_user(c, email='p6b@x.com'))
        resp = self.client.post(reverse('compliance:evidence_upload_v2', args=[item.id]),
                                {'uploaded_file': _f('malware.pdf', EXE), 'notes': ''})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(EvidenceSubmission.objects.count(), before)  # no new spoofed submission

    def test_control_detail_accepts_real_docx(self):
        # DEF-02: a valid .docx must be ACCEPTED and stored (was rejected as 'Unsupported').
        from unittest import mock
        c, control = _company_with_control()
        self.client.force_login(_journey_user(c, email='p6docx@x.com'))
        with mock.patch('compliance.services.process_evidence_pipeline'):  # skip OCR/AI
            resp = self.client.post(reverse('compliance:upload_evidence', args=[control.id]),
                                    {'evidence_file': _f('policy.docx', _docx_bytes())})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Evidence.objects.count(), 1)
        self.assertEqual(Evidence.objects.get().file_type, 'docx')
