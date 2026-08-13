"""
Test suite for CyberTrust KSA (closes the 'no tests' gap, FR Phase-7 QA).
Run: python manage.py test
"""
import io
import os
import tempfile

from django.test import TestCase
from django.urls import reverse

from core.models import AuditLog, Company, EmailVerificationToken, FrameworkDecision, User
from core.forms import CompanyRegistrationForm
from compliance.models import Framework, Domain, Control, CompanyControl


def make_framework_with_controls(code='NCA_ECC', n=3):
    fw, _ = Framework.objects.get_or_create(code=code, defaults={'name': code})
    dom, _ = Domain.objects.get_or_create(framework=fw, name='Governance', defaults={'code': 'GOV'})
    controls = []
    for i in range(1, n + 1):
        controls.append(Control.objects.create(
            framework=fw, domain=dom, control_id=f'{code}-{i}',
            title=f'Control {i}', description='desc', priority='high', evidence_type='policy',
        ))
    return fw, controls


class RegistrationFormTests(TestCase):
    def test_cr_must_be_ten_digits(self):
        form = CompanyRegistrationForm(data={
            'company_name': 'X', 'cr_number': '123', 'sector': 'technology', 'size': 'small',
            'first_name': 'A', 'last_name': 'B', 'email': 'a@b.com',
            'password': 'longenough12', 'nca_scope': True})
        self.assertFalse(form.is_valid())
        self.assertIn('cr_number', form.errors)

    def test_password_min_12(self):
        form = CompanyRegistrationForm(data={
            'company_name': 'X', 'cr_number': '1234567890', 'sector': 'technology', 'size': 'small',
            'first_name': 'A', 'last_name': 'B', 'email': 'a@b.com',
            'password': 'short', 'nca_scope': True})
        self.assertFalse(form.is_valid())
        self.assertIn('password', form.errors)

    def test_requires_a_target(self):
        form = CompanyRegistrationForm(data={
            'company_name': 'X', 'cr_number': '1234567890', 'sector': 'technology', 'size': 'small',
            'first_name': 'A', 'last_name': 'B', 'email': 'a@b.com', 'password': 'longenough12'})
        self.assertFalse(form.is_valid())

    def test_valid_form(self):
        form = CompanyRegistrationForm(data={
            'company_name': 'Acme', 'cr_number': '1234567890', 'sector': 'technology', 'size': 'small',
            'first_name': 'A', 'last_name': 'B', 'email': 'good@x.com',
            'password': 'longenough12', 'nca_scope': True})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['framework_recommendation']['framework_codes'], ['NCA_ECC'])

    def test_recommendation_explains_multiple_frameworks(self):
        form = CompanyRegistrationForm(data={
            'company_name': 'Acme', 'cr_number': '1234567890', 'sector': 'technology', 'size': 'small',
            'first_name': 'A', 'last_name': 'B', 'email': 'good@x.com',
            'password': 'longenough12', 'aramco_supplier': True, 'sabic_supplier': True})
        self.assertTrue(form.is_valid(), form.errors)
        recommendation = form.cleaned_data['framework_recommendation']
        self.assertEqual(recommendation['framework_codes'], ['ARAMCO_SACS002', 'SABIC_CT'])
        self.assertIn('ARAMCO_SACS002', recommendation['rationale'])


class EmailVerificationTests(TestCase):
    def test_verify_email_flow(self):
        user = User.objects.create_user(email='v@x.com', password='longenough12')
        token = EmailVerificationToken.objects.create(user=user, token=EmailVerificationToken.generate())
        resp = self.client.get(reverse('core:verify_email', args=[token.token]))
        self.assertEqual(resp.status_code, 302)
        user.refresh_from_db()
        token.refresh_from_db()
        self.assertTrue(user.email_verified)
        self.assertTrue(token.used)


class MFATests(TestCase):
    def test_totp_verify(self):
        import pyotp
        from core.services import mfa_provisioning_uri, verify_totp
        user = User.objects.create_user(email='m@x.com', password='longenough12')
        mfa_provisioning_uri(user)
        self.assertTrue(verify_totp(user, pyotp.TOTP(user.mfa_secret).now()))
        self.assertFalse(verify_totp(user, '000000'))


class ExtractionTests(TestCase):
    def test_txt_extraction(self):
        from ai_engine.services import process_uploaded_file
        with tempfile.NamedTemporaryFile('w', suffix='.txt', delete=False, encoding='utf-8') as f:
            f.write('Cybersecurity policy approved by management.')
            path = f.name
        out = process_uploaded_file(path, 'txt')
        os.unlink(path)
        self.assertIn('policy', out['text'])

    def test_docx_extraction(self):
        from docx import Document
        from ai_engine.services import process_uploaded_file
        doc = Document()
        doc.add_paragraph('Incident Response Plan v2')
        path = tempfile.mktemp(suffix='.docx')
        doc.save(path)
        out = process_uploaded_file(path, 'docx')
        os.unlink(path)
        self.assertIn('Incident Response Plan', out['text'])

    def test_xlsx_extraction(self):
        from openpyxl import Workbook
        from ai_engine.services import process_uploaded_file
        wb = Workbook()
        wb.active['A1'] = 'Asset'
        wb.active['B1'] = 'Owner'
        path = tempfile.mktemp(suffix='.xlsx')
        wb.save(path)
        out = process_uploaded_file(path, 'xlsx')
        os.unlink(path)
        self.assertIn('Asset', out['text'])


class MonitoringTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(
            name='Co', cr_number='1111111111', sector='technology', size='small',
            contact_email='c@x.com', target_nca=True)
        _, controls = make_framework_with_controls('NCA_ECC', 4)
        for i, c in enumerate(controls):
            CompanyControl.objects.create(
                company=self.company, control=c,
                status='compliant' if i < 2 else 'non_compliant')

    def test_recalculate_score(self):
        from monitoring.services import recalculate_score
        score = recalculate_score(self.company)
        self.assertEqual(score.controls_total, 4)
        self.assertEqual(score.controls_compliant, 2)
        self.assertEqual(score.overall_score, 50.0)

    def test_monthly_report(self):
        from monitoring.services import generate_monthly_report
        report = generate_monthly_report(self.company)
        self.assertEqual(report.report_data['controls_total'], 4)


class ReportingTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(
            name='RepCo', cr_number='2222222222', sector='technology', size='small',
            contact_email='r@x.com', target_nca=True)
        _, controls = make_framework_with_controls('NCA_ECC', 2)
        for c in controls:
            CompanyControl.objects.create(company=self.company, control=c, status='compliant')

    def test_pdf_bytes(self):
        from dashboard.reports import gap_analysis_pdf
        self.assertTrue(gap_analysis_pdf(self.company).startswith(b'%PDF'))

    def test_excel_bytes(self):
        from dashboard.reports import compliance_excel
        self.assertEqual(compliance_excel(self.company)[:2], b'PK')


class ApiTests(TestCase):
    def test_register_and_jwt_login(self):
        make_framework_with_controls('NCA_ECC', 2)
        payload = {
            'company_name': 'ApiCo', 'cr_number': '3333333333', 'sector': 'technology', 'size': 'small',
            'email': 'api@x.com', 'password': 'longenough12', 'first_name': 'A', 'last_name': 'B',
            'nca_scope': True}
        r = self.client.post('/api/v1/register/', data=payload, content_type='application/json')
        self.assertEqual(r.status_code, 201, r.content)
        self.assertIn('access', r.json())
        company = Company.objects.get(cr_number='3333333333')
        self.assertEqual(CompanyControl.objects.filter(company=company).count(), 2)
        decision = FrameworkDecision.objects.get(company=company)
        self.assertEqual(decision.recommended_framework_codes, ['NCA_ECC'])
        self.assertEqual(decision.answers, {'nca_scope': True, 'aramco_supplier': False, 'sabic_supplier': False})
        auth = {'HTTP_AUTHORIZATION': f"Bearer {r.json()['access']}"}
        rc = self.client.get('/api/v1/controls/', **auth)
        self.assertEqual(rc.status_code, 200)
        self.assertEqual(len(rc.json()), 2)

    def test_controls_requires_auth(self):
        self.assertEqual(self.client.get('/api/v1/controls/').status_code, 401)


class AuditLogTests(TestCase):
    def test_post_is_logged(self):
        User.objects.create_user(email='log@x.com', password='longenough12')
        self.client.post(reverse('core:login'),
                         {'username': 'log@x.com', 'password': 'longenough12'})
        self.assertTrue(AuditLog.objects.filter(path=reverse('core:login')).exists())


class PdplTests(TestCase):
    def test_purge_command_runs(self):
        from django.core.management import call_command
        out = io.StringIO()
        call_command('purge_expired_data', stdout=out)
        self.assertIn('Retention purge complete', out.getvalue())
